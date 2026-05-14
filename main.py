"""
FastAPI 后端：上传、聊天修改、撤销、下载接口 + 静态首页。

版本控制：每次 chat 操作前将 current.docx 快照存入 sessions/{id}/history/，
最多保留 MAX_HISTORY 步，/undo 接口弹出最新快照恢复。
"""
import io
import re as _re
import shutil
import uuid
import zipfile
from pathlib import Path
from typing import List

from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.responses import FileResponse, HTMLResponse

import json as _json

from llm_parser import parse_command
from docx_formatter import (
    apply_operations,
    get_used_paragraph_styles,
    get_document_structure,
    get_document_title,
    get_thesis_structure,
)
from schemas import ParsedCommand
from defaults import get_defaults_metadata, collect_ops


app = FastAPI(title="Docx Chat Editor")

WORK_DIR   = Path("./sessions")
BATCH_DIR  = WORK_DIR / "_batch"
STATIC_DIR = Path(__file__).parent / "static"
MAX_HISTORY = 20

WORK_DIR.mkdir(exist_ok=True)
BATCH_DIR.mkdir(exist_ok=True)


# ── 版本历史工具 ──

def _hist_dir(session_dir: Path) -> Path:
    d = session_dir / "history"
    d.mkdir(exist_ok=True)
    return d


def _history_count(session_dir: Path) -> int:
    d = session_dir / "history"
    return len(list(d.glob("*.docx"))) if d.exists() else 0


def _save_snapshot(session_dir: Path):
    """将 current.docx 存入历史；超出上限时丢弃最旧的快照。"""
    hd = _hist_dir(session_dir)
    files = sorted(hd.glob("*.docx"))
    while len(files) >= MAX_HISTORY:
        files[0].unlink()
        files = files[1:]
    next_idx = (int(files[-1].stem) + 1) if files else 0
    shutil.copy2(session_dir / "current.docx", hd / f"{next_idx:06d}.docx")


def _restore_snapshot(session_dir: Path) -> bool:
    """弹出最新快照恢复为 current.docx，返回是否成功。"""
    hd = session_dir / "history"
    if not hd.exists():
        return False
    files = sorted(hd.glob("*.docx"))
    if not files:
        return False
    shutil.copy2(files[-1], session_dir / "current.docx")
    files[-1].unlink()
    return True


# ── 路由 ──

@app.get("/", response_class=HTMLResponse)
async def index():
    return (STATIC_DIR / "index.html").read_text(encoding="utf-8")


@app.post("/upload")
async def upload_doc(file: UploadFile = File(...)):
    if not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "只支持 .docx 文件")

    session_id = uuid.uuid4().hex
    session_dir = WORK_DIR / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    with open(session_dir / "current.docx", "wb") as f:
        shutil.copyfileobj(file.file, f)

    return {"session_id": session_id, "filename": file.filename}


@app.post("/chat")
async def chat(session_id: str = Form(...), message: str = Form(...)):
    session_dir = WORK_DIR / session_id
    doc_path = session_dir / "current.docx"
    if not doc_path.exists():
        raise HTTPException(404, "会话不存在或文档已丢失，请先上传 docx")

    # 1. 提取文档样式 + 大纲 + 推断的论文题目，作为 LLM 上下文
    try:
        available_styles = get_used_paragraph_styles(str(doc_path))
        doc_structure    = get_document_structure(str(doc_path))
        doc_title        = get_document_title(str(doc_path))
    except Exception:
        available_styles = None
        doc_structure    = None
        doc_title        = ""

    # 2. LLM 解析
    try:
        parsed = parse_command(message,
                               available_styles=available_styles,
                               doc_structure=doc_structure,
                               doc_title=doc_title)
    except Exception as e:
        raise HTTPException(500, f"指令解析失败：{e}")

    # 3. 操作前保存快照
    _save_snapshot(session_dir)

    # 4. 应用修改；失败时自动回滚快照
    try:
        apply_operations(str(doc_path), str(doc_path), parsed.operations)
    except Exception as e:
        _restore_snapshot(session_dir)
        raise HTTPException(500, f"文档修改失败：{e}")

    return {
        "explanation":   parsed.explanation,
        "operations":    [op.model_dump(exclude_none=True) for op in parsed.operations],
        "history_count": _history_count(session_dir),
        "distilled":     parsed.distilled,    # 长文本蒸馏结果（短输入时为空）
    }


@app.post("/undo/{session_id}")
async def undo(session_id: str):
    session_dir = WORK_DIR / session_id
    if not (session_dir / "current.docx").exists():
        raise HTTPException(404, "会话不存在")
    if not _restore_snapshot(session_dir):
        raise HTTPException(400, "没有可撤销的操作")
    return {"history_count": _history_count(session_dir)}


@app.get("/download/{session_id}")
async def download(session_id: str):
    doc_path = WORK_DIR / session_id / "current.docx"
    if not doc_path.exists():
        raise HTTPException(404, "文档不存在")
    return FileResponse(
        doc_path,
        filename="modified.docx",
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


@app.get("/structure/{session_id}")
async def get_structure(session_id: str):
    doc_path = WORK_DIR / session_id / "current.docx"
    if not doc_path.exists():
        raise HTTPException(404, "会话不存在或文档已丢失")
    try:
        items = get_thesis_structure(str(doc_path))
    except Exception as e:
        raise HTTPException(500, f"结构识别失败：{e}")
    # Compute summary stats
    from collections import Counter
    counts = Counter(x["type"] for x in items)
    return {
        "items": items,
        "stats": {
            "headings":  sum(v for k, v in counts.items() if k.startswith("heading_")),
            "figures":   counts.get("figure_caption", 0),
            "tables":    counts.get("table_caption",  0),
            "equations": counts.get("equation",       0),
        },
    }


# ── 批量排版（多文档同规则处理）──

def _safe_filename(name: str) -> str:
    """文件名清洗：剔除目录穿越和 Windows 非法字符，保持扩展名。"""
    name = Path(name or "doc.docx").name  # 去掉路径
    # Windows 禁用字符 + 控制字符
    name = _re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)
    name = name.strip(". ")
    return name or "doc.docx"


@app.post("/batch")
async def batch_format(
    files: List[UploadFile] = File(...),
    feature_ids: str = Form("[]"),
    message: str = Form(""),
):
    """对多份 docx 应用同一套排版（默认功能勾选 + 可选自定义指令）。

    feature_ids: JSON 数组字符串，对应 /apply_defaults 的格式
    message:     可选自定义指令；若提供则用首份文档的上下文解析一次，复用到全部文件
    返回：批次 id、每份文件的处理状态；可通过 GET /batch/{id}/download 下载 zip。
    """
    if not files:
        raise HTTPException(400, "至少上传一份 .docx 文件")
    docx_files = [f for f in files if (f.filename or "").lower().endswith(".docx")]
    if not docx_files:
        raise HTTPException(400, "所有上传文件均不是 .docx")

    try:
        feature_id_list = _json.loads(feature_ids)
        if not isinstance(feature_id_list, list):
            raise ValueError("feature_ids 必须是数组")
    except (ValueError, _json.JSONDecodeError) as e:
        raise HTTPException(400, f"feature_ids 格式错误：{e}")

    batch_id = uuid.uuid4().hex
    batch_dir = BATCH_DIR / batch_id
    in_dir    = batch_dir / "in"
    out_dir   = batch_dir / "out"
    in_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)

    # 1) 落盘所有文件（避免 UploadFile 流式状态污染）+ 处理重名
    saved = []   # [(safe_name, in_path)]
    seen  = set()
    for f in docx_files:
        safe = _safe_filename(f.filename or "doc.docx")
        base, ext = (safe[:-5], ".docx") if safe.endswith(".docx") else (safe, ".docx")
        candidate = safe
        i = 1
        while candidate in seen:
            candidate = f"{base}({i}){ext}"
            i += 1
        seen.add(candidate)
        in_path = in_dir / candidate
        with open(in_path, "wb") as fp:
            shutil.copyfileobj(f.file, fp)
        saved.append((candidate, in_path))

    # 2) 自定义指令：仅解析一次（用首份文档的上下文，section: 类目标按首份文档对齐）
    custom_ops_raw = []
    custom_parse_error = None
    if message and message.strip():
        try:
            first_path = saved[0][1]
            try:
                ctx_styles = get_used_paragraph_styles(str(first_path))
                ctx_struct = get_document_structure(str(first_path))
                ctx_title  = get_document_title(str(first_path))
            except Exception:
                ctx_styles = ctx_struct = None
                ctx_title  = ""
            parsed = parse_command(
                message,
                available_styles=ctx_styles,
                doc_structure=ctx_struct,
                doc_title=ctx_title,
            )
            custom_ops_raw = [op.model_dump(exclude_none=True) for op in parsed.operations]
        except Exception as e:
            custom_parse_error = str(e)

    # 3) 逐份处理
    results = []
    for name, in_path in saved:
        if custom_parse_error is not None:
            # 自定义指令解析失败 → 整批失败（用户应当意识到指令问题）
            results.append({"filename": name, "status": "error",
                            "error": f"自定义指令解析失败：{custom_parse_error}"})
            continue

        try:
            # 每份文档单独解析 {THESIS_TITLE} 等占位符
            try:
                thesis_title = get_document_title(str(in_path))
            except Exception:
                thesis_title = ""
            default_ops_raw = collect_ops(feature_id_list, thesis_title=thesis_title)
            all_ops_raw = default_ops_raw + custom_ops_raw
            if not all_ops_raw:
                results.append({"filename": name, "status": "error",
                                "error": "未勾选默认功能且无自定义指令"})
                continue

            # 走与 /chat 同样的 Pydantic 校验
            pc = ParsedCommand.model_validate({"operations": all_ops_raw, "explanation": ""})

            out_path = out_dir / name
            apply_operations(str(in_path), str(out_path), pc.operations)
            results.append({
                "filename": name,
                "status": "ok",
                "ops_count": len(pc.operations),
                "title_used": thesis_title,
            })
        except Exception as e:
            results.append({"filename": name, "status": "error", "error": str(e)})

    # 4) 打包成功项为 zip
    ok_count = sum(1 for r in results if r["status"] == "ok")
    zip_path = batch_dir / "results.zip"
    if ok_count > 0:
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for r in results:
                if r["status"] == "ok":
                    src = out_dir / r["filename"]
                    if src.exists():
                        zf.write(src, arcname=r["filename"])

    return {
        "batch_id":      batch_id,
        "total":         len(results),
        "success_count": ok_count,
        "fail_count":    len(results) - ok_count,
        "results":       results,
        "download_url":  f"/batch/{batch_id}/download" if ok_count > 0 else None,
    }


@app.get("/batch/{batch_id}/download")
async def batch_download(batch_id: str):
    # 防止目录穿越：只接受十六进制 uuid
    if not _re.fullmatch(r"[0-9a-f]{32}", batch_id):
        raise HTTPException(400, "非法 batch_id")
    zip_path = BATCH_DIR / batch_id / "results.zip"
    if not zip_path.exists():
        raise HTTPException(404, "批次不存在或已过期")
    return FileResponse(
        zip_path,
        filename=f"batch_{batch_id[:8]}.zip",
        media_type="application/zip",
    )


# ── 默认排版功能 ──

@app.get("/defaults")
async def list_defaults():
    """返回前端默认排版面板可勾选的全部条目（不含 ops 详情）。"""
    return {"defaults": get_defaults_metadata()}


@app.post("/apply_defaults")
async def apply_defaults(
    session_id: str = Form(...),
    feature_ids: str = Form(...),
):
    """批量应用所选默认排版条目。

    feature_ids: JSON 编码的字符串数组，如 ["page_a4_standard","body_song_xiaosi"]
    """
    session_dir = WORK_DIR / session_id
    doc_path = session_dir / "current.docx"
    if not doc_path.exists():
        raise HTTPException(404, "会话不存在或文档已丢失，请先上传 docx")

    try:
        ids = _json.loads(feature_ids)
        if not isinstance(ids, list):
            raise ValueError("feature_ids 必须是数组")
    except (ValueError, _json.JSONDecodeError) as e:
        raise HTTPException(400, f"feature_ids 格式错误：{e}")

    if not ids:
        raise HTTPException(400, "至少勾选一项排版功能")

    # 推断论文题目，用于 {THESIS_TITLE} 占位符替换
    try:
        thesis_title = get_document_title(str(doc_path))
    except Exception:
        thesis_title = ""

    raw_ops = collect_ops(ids, thesis_title=thesis_title)
    if not raw_ops:
        raise HTTPException(400, "所选条目无可执行操作（可能因占位符无法解析全部跳过）")

    # 复用 ParsedCommand 校验：raw_ops 必须能通过 AnyOperation discriminated union
    try:
        parsed = ParsedCommand.model_validate({"operations": raw_ops, "explanation": ""})
    except Exception as e:
        raise HTTPException(500, f"默认 ops 校验失败：{e}")

    _save_snapshot(session_dir)
    try:
        apply_operations(str(doc_path), str(doc_path), parsed.operations)
    except Exception as e:
        _restore_snapshot(session_dir)
        raise HTTPException(500, f"应用默认排版失败：{e}")

    # 报告：勾选了哪些条目 / 实际生成多少 op / 是否使用了论文题目
    applied_names = [
        d["name"] for d in get_defaults_metadata() if d["id"] in ids
    ]
    note = ""
    if thesis_title and any(
        "{THESIS_TITLE}" in str(_json.dumps(d, ensure_ascii=False))
        for d in [d for d in __import__("defaults").DEFAULTS if d["id"] in ids]
    ):
        note = f"（页眉已自动填入识别出的论文题目：{thesis_title}）"

    return {
        "explanation": (
            f"已应用 {len(applied_names)} 项默认排版，共生成 {len(parsed.operations)} "
            f"个 op 按规范顺序执行。{note}"
        ),
        "applied":      applied_names,
        "operations":   [op.model_dump(exclude_none=True) for op in parsed.operations],
        "history_count": _history_count(session_dir),
    }


# ── 需求文档解析（PDF / DOCX / TXT → 纯文本）──

# OCR 实例：首次使用时延迟加载，避免应用启动时间过长
_OCR_INSTANCE = None

def _get_ocr():
    """单例加载 RapidOCR；首次调用约 1-2 秒。"""
    global _OCR_INSTANCE
    if _OCR_INSTANCE is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR_INSTANCE = RapidOCR()
    return _OCR_INSTANCE


def _extract_pdf_text_layer(data: bytes) -> str:
    """优先 pypdf，其次 PyPDF2 提取 PDF 中的文本图层。扫描件返回空串。"""
    reader = None
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
        except ImportError:
            return ""
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts).strip()


def _extract_pdf_ocr(data: bytes, dpi: int = 200, max_pages: int = 30) -> str:
    """对 PDF 每页渲染成图后逐页 OCR。适用于扫描件或图片型 PDF。

    - dpi:        渲染分辨率，200 在速度/精度间较平衡
    - max_pages:  最多处理页数，防止超大 PDF 阻塞
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise HTTPException(500, "缺少 PDF 渲染依赖：pip install pymupdf")
    try:
        ocr = _get_ocr()
    except ImportError:
        raise HTTPException(500, "缺少 OCR 依赖：pip install rapidocr_onnxruntime")

    pdf = fitz.open(stream=data, filetype="pdf")
    pages = pdf.page_count
    if pages > max_pages:
        pages_to_process = max_pages
    else:
        pages_to_process = pages

    zoom = dpi / 72.0
    mat  = fitz.Matrix(zoom, zoom)

    parts = []
    for i in range(pages_to_process):
        page = pdf.load_page(i)
        pix  = page.get_pixmap(matrix=mat, alpha=False)
        img_bytes = pix.tobytes("png")
        try:
            result, _ = ocr(img_bytes)
        except Exception:
            continue
        if result:
            # result 形如 [[box, text, score], ...]
            page_text = "\n".join(line[1] for line in result if line and len(line) >= 2)
            if page_text.strip():
                parts.append(page_text)
    pdf.close()

    out = "\n\n".join(parts).strip()
    if pages > max_pages:
        out += f"\n\n[文件共 {pages} 页，已 OCR 前 {max_pages} 页]"
    return out


def _extract_pdf(data: bytes):
    """先尝试提取文本图层；若文字过少（疑似扫描件）自动回退 OCR。

    返回 (text, ocr_used)。
    """
    text = _extract_pdf_text_layer(data)
    # 文本图层结果过少（少于 20 个非空白字符）→ 判定为扫描件，走 OCR
    meaningful = sum(1 for ch in text if not ch.isspace())
    if meaningful >= 20:
        return text, False
    return _extract_pdf_ocr(data), True


def _extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    # 兼容表格里的格式要求（如毕业论文模板要求表）
    for t in doc.tables:
        for row in t.rows:
            row_texts = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_texts:
                parts.append(" | ".join(row_texts))
    return "\n".join(parts).strip()


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return data.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


@app.post("/extract_text")
async def extract_text(file: UploadFile = File(...)):
    """从上传的 PDF / DOCX / TXT 中提取纯文本，供前端注入到指令输入框。

    PDF：先尝试文字图层，无文字（扫描件）自动走 OCR。
    """
    name = (file.filename or "").lower()
    data = await file.read()
    if not data:
        raise HTTPException(400, "文件为空")

    ocr_used = False
    try:
        if name.endswith(".txt"):
            text = _extract_txt(data)
        elif name.endswith(".docx"):
            text = _extract_docx(data)
        elif name.endswith(".pdf"):
            text, ocr_used = _extract_pdf(data)
        else:
            raise HTTPException(400, "仅支持 .pdf / .docx / .txt 文件")
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"文件解析失败：{e}")

    if not text:
        raise HTTPException(400, "未从文件中提取到任何文字")

    return {
        "text": text,
        "filename": file.filename,
        "chars": len(text),
        "ocr_used": ocr_used,
    }


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
