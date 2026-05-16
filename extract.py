"""
文件文本提取工具（PDF / DOCX / TXT → 纯文本）。

被 main.py（需求文档解析）与 knowledge_base.py（国标 PDF 入库）共用。
PDF：优先文字图层，扫描件自动回退 OCR。
"""
import io


# OCR 实例：首次使用时延迟加载，避免应用启动时间过长
_OCR_INSTANCE = None


def _get_ocr():
    """单例加载 RapidOCR；首次调用约 1-2 秒。"""
    global _OCR_INSTANCE
    if _OCR_INSTANCE is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR_INSTANCE = RapidOCR()
    return _OCR_INSTANCE


def extract_pdf_text_layer(data: bytes) -> str:
    """优先 pypdf，其次 PyPDF2 提取 PDF 中的文本图层。扫描件返回空串。"""
    reader = None
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
    except ImportError:
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(io.BytesIO(data))
        except ImportError:
            return ""
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            continue
    return "\n".join(parts).strip()


def extract_pdf_ocr(data: bytes, dpi: int = 200, max_pages: int = 30) -> str:
    """对 PDF 每页渲染成图后逐页 OCR。适用于扫描件或图片型 PDF。

    - dpi:        渲染分辨率，200 在速度/精度间较平衡
    - max_pages:  最多处理页数，防止超大 PDF 阻塞
    """
    try:
        import fitz  # PyMuPDF
    except ImportError:
        raise RuntimeError("缺少 PDF 渲染依赖：pip install pymupdf")
    try:
        ocr = _get_ocr()
    except ImportError:
        raise RuntimeError("缺少 OCR 依赖：pip install rapidocr_onnxruntime")

    pdf = fitz.open(stream=data, filetype="pdf")
    pages = pdf.page_count
    pages_to_process = min(pages, max_pages)

    zoom = dpi / 72.0
    mat  = fitz.Matrix(zoom, zoom)

    parts = []
    for i in range(pages_to_process):
        page = pdf.load_page(i)
        pix  = page.get_pixmap(matrix=mat, alpha=False)
        img_bytes = pix.tobytes("png")
        try:
            result, _ = ocr(img_bytes)
        except Exception:
            continue
        if result:
            page_text = "\n".join(line[1] for line in result if line and len(line) >= 2)
            if page_text.strip():
                parts.append(page_text)
    pdf.close()

    out = "\n\n".join(parts).strip()
    if pages > max_pages:
        out += f"\n\n[文件共 {pages} 页，已 OCR 前 {max_pages} 页]"
    return out


def extract_pdf(data: bytes):
    """先尝试提取文本图层；若文字过少（疑似扫描件）自动回退 OCR。

    返回 (text, ocr_used)。
    """
    text = extract_pdf_text_layer(data)
    meaningful = sum(1 for ch in text if not ch.isspace())
    if meaningful >= 20:
        return text, False
    return extract_pdf_ocr(data), True


def extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text)
    for t in doc.tables:
        for row in t.rows:
            row_texts = [c.text.strip() for c in row.cells if c.text.strip()]
            if row_texts:
                parts.append(" | ".join(row_texts))
    return "\n".join(parts).strip()


def extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "gbk", "gb18030"):
        try:
            return data.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace").strip()


def extract_any(filename: str, data: bytes):
    """按扩展名分派提取。返回 (text, ocr_used)。不支持的类型抛 ValueError。"""
    name = (filename or "").lower()
    if name.endswith(".txt"):
        return extract_txt(data), False
    if name.endswith(".docx"):
        return extract_docx(data), False
    if name.endswith(".pdf"):
        return extract_pdf(data)
    raise ValueError("仅支持 .pdf / .docx / .txt 文件")
