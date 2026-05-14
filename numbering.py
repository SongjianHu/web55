"""
图、表、公式按章节自动编号，支持 Word 书签交叉引用。

编号规则：X-Y（X=章节号对应 Heading 1 序号，Y=章内流水号）
样式：
  图题     → 图 X-Y [说明]   （居中，图片下方）
  表题     → 表 X-Y [说明]   （居中，表格上方）
  公式编号 → (X-Y)            （居中）
书签命名：fig_X_Y / tab_X_Y / eq_X_Y
"""
import re
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE


CAPTION_STYLE_MAP = {"图": "图题", "表": "表题", "公式": "公式编号"}
BOOKMARK_PREFIX   = {"图": "fig",  "表": "tab",  "公式": "eq"}

# 内容识别：以 图/表 + 数字（含连字符变体、点）开头，后跟空白或行尾
# 公式：整段为 (数字) 形式
_DASH = r"\-−–—－‐"   # 普通连字符 / 减号 / en-dash / em-dash / 全角连字符 / hyphen
_CONTENT_PAT = {
    "图":   re.compile(rf"^图\s*(\d[\d{_DASH}.]*|X[{_DASH}]Y)(?=\s|$)"),
    "表":   re.compile(rf"^表\s*(\d[\d{_DASH}.]*|X[{_DASH}]Y)(?=\s|$)"),
    "公式": re.compile(rf"^\s*\((\d[\d{_DASH}.]*|X[{_DASH}]Y)\)\s*$"),
}
# Word 内置题注样式名（中/英文）
_BUILTIN_CAPTION_STYLES = {"题注", "Caption"}


# -------- 内部工具 --------

def _get_heading_level(style_name: str):
    """返回 Heading N 层级；兼容 'Heading 1' / '标题 1' / '标题1' 等。"""
    if not style_name:
        return None
    for pat in (r"^Heading\s*(\d)$", r"^标题\s*(\d)$"):
        m = re.match(pat, style_name)
        if m:
            return int(m.group(1))
    return None


def _ensure_caption_styles(doc):
    """创建缺失的图题/表题/公式编号样式。"""
    for name in ["图题", "表题", "公式编号"]:
        try:
            doc.styles[name]
        except KeyError:
            s = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            s.base_style = doc.styles["Normal"]
            s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            s.paragraph_format.space_before = Pt(6)
            s.paragraph_format.space_after = Pt(6)


def _max_bookmark_id(doc) -> int:
    mx = 0
    for bm in doc.element.body.iter(qn("w:bookmarkStart")):
        try:
            mx = max(mx, int(bm.get(qn("w:id"), 0)))
        except (ValueError, TypeError):
            pass
    return mx


def _extract_desc(text: str, prefix: str) -> str:
    """从题注文字中提取说明部分，兼容多种已有格式。"""
    text = text.strip()
    # 已编号格式："图 1-1 说明" / "图1-1 说明" / "图1.1 说明" / "图1 说明"
    m = re.match(rf"^{re.escape(prefix)}\s*\d[\d{_DASH}.]*\s*", text)
    if m:
        return text[m.end():].strip()
    # 占位符格式："图 X-Y 说明"
    m = re.match(rf"^{re.escape(prefix)}\s*X[{_DASH}]Y\s*", text, re.IGNORECASE)
    if m:
        return text[m.end():].strip()
    # 仅前缀后跟空格："图 说明"
    m = re.match(rf"^{re.escape(prefix)}\s+", text)
    if m:
        return text[m.end():].strip()
    return ""


# -------- XML 直接访问（绕过 doc.paragraphs，覆盖表格内段落）--------

def _p_text(p_elem) -> str:
    """从 w:p 元素聚合所有 w:t 文字。"""
    return "".join(t.text or "" for t in p_elem.iter(qn("w:t")))


def _p_style_id(p_elem):
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        return None
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        return None
    return pStyle.get(qn("w:val"))


def _style_id_for(doc, style_name: str) -> str:
    try:
        return doc.styles[style_name].style_id
    except (KeyError, AttributeError):
        return style_name


def _set_p_style(p_elem, style_id: str):
    """设置 w:pStyle/@w:val。"""
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        pPr = OxmlElement("w:pPr")
        p_elem.insert(0, pPr)
    pStyle = pPr.find(qn("w:pStyle"))
    if pStyle is None:
        pStyle = OxmlElement("w:pStyle")
        pPr.insert(0, pStyle)
    pStyle.set(qn("w:val"), style_id)


def _detect_cap_type(text: str):
    """通过文字识别题注类型。短段落且匹配特定前缀模式才返回，避免误判正文。"""
    if not text:
        return None
    text_stripped = text.strip()
    if not text_stripped or len(text_stripped) > 120:
        return None
    for ctype, pat in _CONTENT_PAT.items():
        if pat.match(text_stripped):
            return ctype
    return None


def _write_caption(p_elem, cap_type: str, num: str, desc: str, bm_name: str, bm_id: int):
    """覆写段落内容为带书签的题注文字（保留 pPr 中的样式/对齐设置）。"""
    pPr = p_elem.find(qn("w:pPr"))
    for child in list(p_elem):
        if child is not pPr:
            p_elem.remove(child)

    if cap_type == "公式":
        label = f"({num})"
    else:
        label = f"{cap_type} {num}" + (f" {desc}" if desc else "")

    bm_start = OxmlElement("w:bookmarkStart")
    bm_start.set(qn("w:id"), str(bm_id))
    bm_start.set(qn("w:name"), bm_name)
    p_elem.append(bm_start)

    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = label
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    p_elem.append(r)

    bm_end = OxmlElement("w:bookmarkEnd")
    bm_end.set(qn("w:id"), str(bm_id))
    p_elem.append(bm_end)


# -------- 公开 API --------

def renumber_all(doc):
    """扫描全文，对图题/表题/公式编号段落按章节重新编号并写入书签。

    识别逻辑（优先级从高到低）：
    1. 段落样式为 "图题"/"表题"/"公式编号"
    2. 段落样式为 Word 内置 "题注"/"Caption"，按文字推断类型
    3. 普通段落但文字符合题注模式（自动升级样式）

    覆盖范围：正文段落 + 表格单元格段落（按文档顺序遍历）。
    """
    _ensure_caption_styles(doc)
    style_to_type = {"图题": "图", "表题": "表", "公式编号": "公式"}

    body = doc.element.body
    # style_id → name 映射，避免 O(N) 多次查样式
    id_to_name = {}
    name_to_style = {}
    for s in doc.styles:
        try:
            id_to_name[s.style_id] = s.name
            name_to_style[s.name] = s
        except Exception:
            pass

    def get_sname(p_elem):
        sid = _p_style_id(p_elem)
        if sid is None:
            return "Normal"
        return id_to_name.get(sid, sid)

    def get_heading_level(sname):
        """沿 base_style 链向上查找 Heading N 层级，覆盖派生自 Heading 1 的自定义样式。"""
        seen = set()
        cur = sname
        while cur and cur not in seen:
            seen.add(cur)
            lvl = _get_heading_level(cur)
            if lvl is not None:
                return lvl
            style = name_to_style.get(cur)
            if style is None:
                return None
            try:
                base = style.base_style
                cur = base.name if base is not None else None
            except Exception:
                return None
        return None

    # 1) 检测是否存在 H1（仅 body 顶层段落参与判断）
    has_chapters = False
    for child in body:
        if child.tag == qn("w:p"):
            if get_heading_level(get_sname(child)) == 1:
                has_chapters = True
                break

    chapter = 0 if has_chapters else 1
    counts  = {"图": 0, "表": 0, "公式": 0}
    bm_id   = _max_bookmark_id(doc) + 1

    # 2) 按文档顺序遍历所有 w:p（含表格内）
    for p_elem in body.iter(qn("w:p")):
        sname = get_sname(p_elem)

        # 仅顶层段落参与章节计数
        if p_elem.getparent() is body:
            level = get_heading_level(sname)
            if level == 1:
                prev = chapter
                chapter += 1
                # 仅当从已有章节切换到下一章时才重置；从 0→1 保留之前累计的题注计数
                if prev > 0:
                    counts = {"图": 0, "表": 0, "公式": 0}
                continue

        # --- 识别题注类型 ---
        text = _p_text(p_elem)
        cap_type = style_to_type.get(sname)

        if cap_type is None and sname in _BUILTIN_CAPTION_STYLES:
            cap_type = _detect_cap_type(text)

        if cap_type is None:
            # 内容识别：跳过含图片的段落，避免重写时删除内嵌图片
            has_drawing = next(p_elem.iter(qn("w:drawing")), None) is not None
            if not has_drawing:
                detected = _detect_cap_type(text)
                if detected is not None:
                    cap_type = detected
                    # 升级为专属样式
                    target = CAPTION_STYLE_MAP[detected]
                    try:
                        _set_p_style(p_elem, name_to_style[target].style_id)
                    except KeyError:
                        pass

        if cap_type is None:
            continue

        # 章前题注归入第 1 章命名空间（chapter==0 → ch=1）
        ch = chapter if chapter > 0 else 1
        counts[cap_type] += 1
        num     = f"{ch}-{counts[cap_type]}"
        bm_name = f"{BOOKMARK_PREFIX[cap_type]}_{ch}_{counts[cap_type]}"
        desc    = _extract_desc(text, cap_type)
        _write_caption(p_elem, cap_type, num, desc, bm_name, bm_id)
        bm_id += 1


def insert_caption(doc, caption_type: str, after_text: str = None, description: str = None):
    """
    插入一个新的题注段落（占位符），随后由 renumber_all 更新编号。

    caption_type : "图" | "表" | "公式"
    after_text   : 在含此文字的段落之后插入；None 则追加到文档末尾
    description  : 题注说明文字；None 则使用占位符
    """
    _ensure_caption_styles(doc)
    style_name  = CAPTION_STYLE_MAP[caption_type]
    desc        = description or "[请填写说明]"
    placeholder = f"{caption_type} X-Y {desc}" if caption_type != "公式" else "(X-Y)"

    # 在含 after_text 的段落（含表格内段落）之后插入
    ref_p_elem = None
    if after_text:
        for p_elem in doc.element.body.iter(qn("w:p")):
            if after_text in _p_text(p_elem):
                ref_p_elem = p_elem
                break

    if ref_p_elem is not None:
        new_p = OxmlElement("w:p")
        pPr   = OxmlElement("w:pPr")
        pSty  = OxmlElement("w:pStyle")
        pSty.set(qn("w:val"), _style_id_for(doc, style_name))
        pPr.append(pSty)
        new_p.append(pPr)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = placeholder
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        new_p.append(r)
        ref_p_elem.addnext(new_p)
    else:
        # 追加到文档末尾
        p = doc.add_paragraph(style=style_name)
        p.add_run(placeholder)


def insert_cross_ref(doc, bookmark_name: str, display_text: str,
                     in_paragraph_containing: str = None):
    """
    在指定段落末尾追加 Word REF 域（交叉引用）。

    bookmark_name           : 书签名，如 fig_1_1 / tab_2_3 / eq_1_2
    display_text            : 缓存显示文字，如 图1-1（供 docx-preview 显示）
    in_paragraph_containing : 定位段落：含此文字的第一个段落（含表格内段落）
    """
    target = None
    if in_paragraph_containing:
        for p_elem in doc.element.body.iter(qn("w:p")):
            if in_paragraph_containing in _p_text(p_elem):
                target = p_elem
                break
    if target is None:
        return

    def _run(*children):
        r = OxmlElement("w:r")
        for c in children:
            r.append(c)
        target.append(r)

    fc_begin = OxmlElement("w:fldChar"); fc_begin.set(qn("w:fldCharType"), "begin")
    _run(fc_begin)

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" REF {bookmark_name} \\h "
    _run(instr)

    fc_sep = OxmlElement("w:fldChar"); fc_sep.set(qn("w:fldCharType"), "separate")
    _run(fc_sep)

    t = OxmlElement("w:t"); t.text = display_text
    _run(t)

    fc_end = OxmlElement("w:fldChar"); fc_end.set(qn("w:fldCharType"), "end")
    _run(fc_end)
