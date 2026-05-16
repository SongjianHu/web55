"""
检查模式：对上传的学位论文做合规性检查。

四类检查（可按需选择子集）：
  citation          引用对应性：正文[n] ↔ 参考文献表 ↔ 文献综述 三方交叉核对
  reference_format  参考文献格式：按 GB/T 7714—2015 顺序编码制启发式校验每条著录
  structure         结构完整性：必备章节、标题层级跳级、图表公式编号连续性
  compliance        格式合规：页面/字体/字号/行距、摘要字数、关键词个数、中英文摘要齐备

返回结构：
  {
    "summary": {"ok": n, "warn": n, "error": n},
    "categories": [
      {"id","title","items":[{"level":"ok|warn|error","msg","detail"?}, ...]}
    ]
  }
设计原则：学校模板差异大 → 格式类多用 warn（提示）而非 error，避免误报。
"""
import re

from docx import Document
from docx.oxml.ns import qn

from docx_formatter import get_thesis_structure, get_document_structure
from extract import extract_any

ALL_CATEGORIES = ["citation", "reference_format", "structure", "compliance"]

# 文献类型/载体标识
_TYPE_TAG_RE = re.compile(r"\[[A-Z]{1,3}(?:/[A-Z]{2})?\]")
# 正文引用标记 [1] [1-3] [1,2] [1，3]（排除参考文献条目本身）
_CITE_RE = re.compile(r"\[(\d{1,3}(?:\s*[-–,，]\s*\d{1,3})*)\]")
_YEAR_RE = re.compile(r"(19|20)\d{2}")
_REF_HEAD_RE = re.compile(r"^\s*(参考文献|REFERENCES?|Bibliography)\s*$", re.I)
_REF_ENTRY_RE = re.compile(r"^\s*\[?(\d{1,3})\]?[\.、]?\s*\S")

# 章节编号前缀：第1章 / 第一章 / 第 一 章 / 第一篇 / Chapter 1 / 1 / 1. / 1.1 等
_CHAP_PREFIX_RE = re.compile(
    r"^\s*(?:"
    r"第\s*[0-9一二三四五六七八九十百零两]+\s*[章篇部节]"   # 第1章/第一章/第一篇/第一节
    r"|chapter\s*[0-9ivxlcdm]+"                            # Chapter 1 / Chapter IV
    r"|[0-9]+(?:\s*[.\.][0-9]+)*"                          # 1 / 1. / 1.1 / 1.1.1
    r")\s*[\.．、,，:：\-—–　]*\s*",
    re.I,
)
# 引言/绪论 与 结论 关键词（去掉编号前缀后，出现在标题开头一小段内才算）
_INTRO_KW_RE = re.compile(r"^(绪\s*论|引\s*言|前\s*言|序\s*言|导\s*论|绪\s*言)")
_CONCL_KW_RE = re.compile(
    r"^(结\s*论|结\s*语|总\s*结|结\s*束\s*语|全\s*文\s*总\s*结"
    r"|结论与展望|总结与展望|结论与建议|工作总结)"
)


def _strip_chapter_prefix(text: str) -> str:
    """去掉标题开头的『第N章/Chapter N/1.1』等编号前缀，便于关键词识别。"""
    return _CHAP_PREFIX_RE.sub("", (text or "").strip(), count=1).strip()


# ============================================================
# 文档读取辅助
# ============================================================

def _doc_paragraphs(path: str):
    doc = Document(path)
    return doc, [p for p in doc.paragraphs]


def _full_text(paragraphs) -> str:
    return "\n".join(p.text for p in paragraphs if p.text and p.text.strip())


def _find_ref_section_index(paragraphs) -> int:
    """返回'参考文献'标题段落的下标；找不到返回 -1。"""
    for i, p in enumerate(paragraphs):
        if _REF_HEAD_RE.match(p.text.strip()):
            return i
    return -1


def _parse_reference_entries(paragraphs, ref_idx: int):
    """从参考文献标题之后解析条目。返回 [(num:int, text:str), ...]。

    遇到下一个一级标题（致谢/附录等）或文字明显不再是著录时停止。
    """
    entries = []
    if ref_idx < 0:
        return entries
    cur_num, cur_buf = None, []
    for p in paragraphs[ref_idx + 1:]:
        t = p.text.strip()
        if not t:
            continue
        # 到达后续章节标题（致谢/附录/攻读学位成果等）停止
        if re.match(r"^(致\s*谢|附\s*录|攻读|个人简历|索\s*引|ACKNOWLEDGE)", t, re.I):
            break
        m = re.match(r"^\[?(\d{1,3})\]?[\.、]?\s*(.+)$", t)
        if m and (cur_num is None or int(m.group(1)) == cur_num + 1
                  or int(m.group(1)) == 1):
            if cur_num is not None:
                entries.append((cur_num, " ".join(cur_buf).strip()))
            cur_num = int(m.group(1))
            cur_buf = [m.group(2)]
        else:
            if cur_num is not None:
                cur_buf.append(t)
    if cur_num is not None:
        entries.append((cur_num, " ".join(cur_buf).strip()))
    return entries


def _expand_cites(token: str):
    """'1-3,5' → {1,2,3,5}。"""
    nums = set()
    for part in re.split(r"[,，]", token):
        part = part.strip()
        m = re.match(r"^(\d{1,3})\s*[-–]\s*(\d{1,3})$", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a <= b and b - a < 200:
                nums.update(range(a, b + 1))
        elif part.isdigit():
            nums.add(int(part))
    return nums


def _collect_citations(text: str):
    cited = set()
    for m in _CITE_RE.finditer(text):
        cited |= _expand_cites(m.group(1))
    return cited


def _norm(s: str) -> str:
    return re.sub(r"[\s\.,;:：，。；、\[\]()（）]+", "", s).lower()


# ============================================================
# 检查 1：引用对应性
# ============================================================

def _check_citation(thesis_paragraphs, review_text: str | None):
    items = []
    ref_idx = _find_ref_section_index(thesis_paragraphs)
    if ref_idx < 0:
        items.append({"level": "error", "msg": "未找到「参考文献」章节，无法核对引用对应性"})
        return items

    body_text = _full_text(thesis_paragraphs[:ref_idx])
    cited = _collect_citations(body_text)
    entries = _parse_reference_entries(thesis_paragraphs, ref_idx)
    listed = {n for n, _ in entries}

    if not cited:
        items.append({"level": "warn", "msg": "正文中未检出 [n] 形式的引用标记（可能用了著者-出版年制或上标域）"})
    if not listed:
        items.append({"level": "error", "msg": "参考文献表为空或无法解析编号"})
        return items

    missing = sorted(cited - listed)
    uncited = sorted(listed - cited)

    if missing:
        items.append({"level": "error",
                       "msg": f"正文引用了 {len(missing)} 个编号但参考文献表中没有：{missing}"})
    if uncited:
        items.append({"level": "warn",
                       "msg": f"参考文献表列出但正文未引用 {len(uncited)} 条：{uncited}"})
    # 编号连续性
    if listed:
        expected = set(range(1, max(listed) + 1))
        gaps = sorted(expected - listed)
        if gaps:
            items.append({"level": "warn", "msg": f"参考文献编号不连续，缺号：{gaps}"})
    if not missing and not uncited:
        items.append({"level": "ok",
                       "msg": f"正文引用与参考文献表对应一致（共 {len(listed)} 条，引用 {len(cited)} 处）"})

    # 文献综述 ↔ 参考文献（启发式：著录文字归一化后包含匹配）
    if review_text:
        review_cites = _collect_citations(review_text)
        ref_norm = [(n, _norm(t)) for n, t in entries]
        review_lines = [ln.strip() for ln in review_text.splitlines() if len(ln.strip()) > 25]
        review_refish = [ln for ln in review_lines if _TYPE_TAG_RE.search(ln) or _YEAR_RE.search(ln)]
        unmatched = []
        for ln in review_refish[:200]:
            nl = _norm(ln)
            if len(nl) < 12:
                continue
            hit = any(nl[:40] and (nl[:40] in rn or rn[:40] in nl) for _, rn in ref_norm)
            if not hit:
                unmatched.append(ln[:60])
        if review_cites:
            items.append({"level": "ok" if review_cites <= listed else "warn",
                           "msg": f"文献综述检出 {len(review_cites)} 个引用编号；"
                                  + ("均能在论文参考文献表编号范围内"
                                     if review_cites <= listed
                                     else f"超出范围：{sorted(review_cites - listed)}")})
        if unmatched:
            items.append({"level": "warn",
                           "msg": f"文献综述中 {len(unmatched)} 条疑似文献未在论文参考文献表中找到对应（启发式匹配，需人工复核）",
                           "detail": "\n".join(f"• {u}" for u in unmatched[:15])})
        else:
            items.append({"level": "ok", "msg": "文献综述涉及文献均可在论文参考文献表中找到对应（启发式）"})

    return items


# ============================================================
# 检查 2：参考文献格式（GB/T 7714—2015）
# ============================================================

def _check_reference_format(thesis_paragraphs):
    items = []
    ref_idx = _find_ref_section_index(thesis_paragraphs)
    if ref_idx < 0:
        items.append({"level": "error", "msg": "未找到「参考文献」章节"})
        return items
    entries = _parse_reference_entries(thesis_paragraphs, ref_idx)
    if not entries:
        items.append({"level": "error", "msg": "参考文献表无法解析出条目"})
        return items

    problems = []
    for n, txt in entries:
        reasons = []
        if not _TYPE_TAG_RE.search(txt):
            reasons.append("缺文献类型标识（如 [M]/[J]/[D]/[C]/[EB/OL]）")
        if not _YEAR_RE.search(txt):
            reasons.append("缺出版年份")
        if not txt.rstrip().endswith("."):
            reasons.append("条目未以英文句点“.”结束")
        if "．" in txt or "，" in txt and "[" not in txt[:3]:
            pass  # 中文标点可接受，不强制
        if len(txt) < 12:
            reasons.append("内容过短，著录项可能不完整")
        if reasons:
            problems.append((n, txt[:70], reasons))

    items.append({"level": "ok",
                   "msg": f"共解析参考文献 {len(entries)} 条"})
    if problems:
        detail = "\n".join(
            f"[{n}] {t}…  →  {'；'.join(rs)}" for n, t, rs in problems[:25]
        )
        lvl = "error" if len(problems) > len(entries) * 0.4 else "warn"
        items.append({"level": lvl,
                       "msg": f"{len(problems)} 条参考文献著录可能不符合 GB/T 7714 规范",
                       "detail": detail})
    else:
        items.append({"level": "ok", "msg": "各条参考文献著录项基本完整（含类型标识、年份、句点结尾）"})
    return items


# ============================================================
# 检查 3：结构完整性
# ============================================================

def _check_structure(thesis_path: str):
    items = []
    try:
        sections = get_thesis_structure(thesis_path)
    except Exception as e:
        items.append({"level": "error", "msg": f"结构识别失败：{e}"})
        return items

    flat_types = {it["type"] for sec in sections.values() for it in sec}

    required = [
        ("abstract_h",      "中文摘要"),
        ("keywords_h",      "关键词"),
        ("en_abstract_h",   "英文摘要 Abstract"),
        ("toc_h",           "目录"),
        ("references",      "参考文献"),
        ("acknowledgment",  "致谢"),
    ]
    missing = [name for key, name in required if key not in flat_types]
    if missing:
        items.append({"level": "warn", "msg": "缺少常见必备部分：" + "、".join(missing)})
    else:
        items.append({"level": "ok", "msg": "前置/后置必备部分齐全（摘要/关键词/英文摘要/目录/参考文献/致谢）"})

    # 引言/绪论 与 结论 识别：
    #   先看 get_thesis_structure 的精确分类（intro_h1/conclusion_h1），
    #   再回退扫描正文章节标题——绪论/结论常被写进章节编号里
    #   （如「第1章 绪论」「第一章 绪论」「1 绪论」「Chapter 1 Introduction」），
    #   去掉编号前缀后再按关键词匹配，避免漏判。
    outline = get_document_structure(thesis_path)
    body_items = sections.get("body", [])
    has_intro = any(it["type"] == "intro_h1" for it in body_items)
    has_conclusion = any(it["type"] == "conclusion_h1" for it in body_items)

    if not (has_intro and has_conclusion):
        cand_texts = [it.get("text", "") for it in body_items]
        cand_texts += [h["text"] for h in outline if h.get("level", 9) <= 1]
        intro_hit = concl_hit = ""
        for raw in cand_texts:
            core = _strip_chapter_prefix(raw)
            if not core:
                continue
            if not has_intro and _INTRO_KW_RE.match(core):
                has_intro, intro_hit = True, raw.strip()
            if not has_conclusion and _CONCL_KW_RE.match(core):
                has_conclusion, concl_hit = True, raw.strip()

    if has_intro:
        items.append({"level": "ok", "msg": "已识别到「引言/绪论」章节"})
    else:
        items.append({"level": "warn", "msg": "未识别到「引言/绪论」章节"})
    if has_conclusion:
        items.append({"level": "ok", "msg": "已识别到「结论」章节"})
    else:
        items.append({"level": "warn", "msg": "未识别到「结论」章节"})

    # 标题层级跳级（复用上面已取的 outline）
    gaps = []
    prev = 0
    for h in outline:
        lv = h["level"]
        if prev and lv > prev + 1:
            gaps.append(f"“{h['text'][:30]}” 由 {prev} 级直接跳到 {lv} 级")
        prev = lv
    if gaps:
        items.append({"level": "warn", "msg": f"标题层级存在跳级 {len(gaps)} 处",
                       "detail": "\n".join("• " + g for g in gaps[:15])})
    elif outline:
        items.append({"level": "ok", "msg": f"标题层级连续无跳级（共 {len(outline)} 个标题）"})

    # 图/表/公式编号连续性
    doc, paras = _doc_paragraphs(thesis_path)
    full = _full_text(paras)
    for label, pat in (("图", r"图\s*(\d+)\s*[-–.]\s*(\d+)"),
                       ("表", r"表\s*(\d+)\s*[-–.]\s*(\d+)"),
                       ("公式", r"[（(]\s*(\d+)\s*[-–.]\s*(\d+)\s*[)）]")):
        seq = {}
        for m in re.finditer(pat, full):
            ch, idx = int(m.group(1)), int(m.group(2))
            seq.setdefault(ch, set()).add(idx)
        broken = []
        for ch, idxs in sorted(seq.items()):
            exp = set(range(1, max(idxs) + 1))
            miss = sorted(exp - idxs)
            if miss:
                broken.append(f"第{ch}章{label}缺号 {miss}")
        if seq:
            if broken:
                items.append({"level": "warn",
                               "msg": f"{label}编号不连续：" + "；".join(broken[:8])})
            else:
                items.append({"level": "ok", "msg": f"{label}编号按章连续"})
    return items


# ============================================================
# 检查 4：格式合规
# ============================================================

_CJK_RE = re.compile(r"[一-鿿]")
_LAT_RE = re.compile(r"[A-Za-z]")


def _run_fonts(run):
    """返回 (中文字体 eastAsia, 西文字体 ascii)；未显式设置则为 None。

    关键：中文/西文分开取——很多论文模板正文是「中文宋体 + 西文 Times New
    Roman」，旧实现 eastAsia 缺失时回退取 ascii，会把西文 TNR 误当成正文主字体。
    """
    ea = asc = None
    try:
        rPr = run._element.rPr
        if rPr is not None:
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is not None:
                ea = rFonts.get(qn("w:eastAsia"))
                asc = rFonts.get(qn("w:ascii"))
    except Exception:
        pass
    if not asc:
        asc = run.font.name  # python-docx 的 .name 实质是 ascii
    return (ea or None), (asc or None)


def _style_east_asia_font(doc, style_names):
    """从 styles.xml 取某些段落样式的 eastAsia 字体（运行级未显式设置时的回退）。"""
    try:
        styles_el = doc.styles.element
        for st in styles_el.findall(qn("w:style")):
            nm = st.find(qn("w:name"))
            sid = st.get(qn("w:styleId")) or ""
            label = (nm.get(qn("w:val")) if nm is not None else "") or sid
            if label not in style_names and sid not in style_names:
                continue
            rPr = st.find(qn("w:rPr"))
            if rPr is None:
                continue
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is not None:
                ea = rFonts.get(qn("w:eastAsia"))
                if ea:
                    return ea
    except Exception:
        pass
    return None


def _is_songti(name: str) -> bool:
    return bool(name) and any(k in name for k in ("宋体", "SimSun", "Song", "STSong", "NSimSun"))


def _check_compliance(thesis_path: str):
    items = []
    doc = Document(thesis_path)

    # 页面与页边距
    try:
        sec = doc.sections[0]
        w_cm = round(sec.page_width.cm, 1)
        h_cm = round(sec.page_height.cm, 1)
        if {w_cm, h_cm} == {21.0, 29.7}:
            items.append({"level": "ok", "msg": f"纸张为 A4（{w_cm}×{h_cm} cm）"})
        else:
            items.append({"level": "warn", "msg": f"纸张非标准 A4：{w_cm}×{h_cm} cm（A4 应为 21×29.7）"})
        ml, mr = round(sec.left_margin.cm, 2), round(sec.right_margin.cm, 2)
        mt, mb = round(sec.top_margin.cm, 2), round(sec.bottom_margin.cm, 2)
        warn_m = []
        if not 2.0 <= ml <= 3.5: warn_m.append(f"左 {ml}cm")
        if not 1.8 <= mr <= 3.0: warn_m.append(f"右 {mr}cm")
        if not 2.0 <= mt <= 3.5: warn_m.append(f"上 {mt}cm")
        if not 2.0 <= mb <= 3.5: warn_m.append(f"下 {mb}cm")
        if warn_m:
            items.append({"level": "warn", "msg": "页边距偏离常见范围：" + "、".join(warn_m)})
        else:
            items.append({"level": "ok", "msg": f"页边距合理（上{mt} 下{mb} 左{ml} 右{mr} cm）"})
    except Exception as e:
        items.append({"level": "warn", "msg": f"页面信息读取失败：{e}"})

    # 正文字体/字号（取样式名含“正文/Normal”的段落）
    # 中文字体与西文字体分别统计：论文规范一般是「中文宋体 + 西文 Times New Roman」，
    # 二者都正确才算合规，不能用西文 TNR 去否定中文字体。
    cjk_fonts, west_fonts, body_sizes = {}, {}, {}
    BODY_STYLE_NAMES = ("Normal", "正文", "Body Text")
    for p in doc.paragraphs:
        sname = ""
        try:
            sname = p.style.name or ""
        except Exception:
            pass
        if sname in BODY_STYLE_NAMES and p.text.strip():
            for r in p.runs:
                if not r.text.strip():
                    continue
                ea, asc = _run_fonts(r)
                if _CJK_RE.search(r.text) and ea:
                    cjk_fonts[ea] = cjk_fonts.get(ea, 0) + 1
                if _LAT_RE.search(r.text) and asc:
                    west_fonts[asc] = west_fonts.get(asc, 0) + 1
                if r.font.size:
                    sz = round(r.font.size.pt, 1)
                    body_sizes[sz] = body_sizes.get(sz, 0) + 1

    # 中文字体：优先 run 级显式 eastAsia；都没有则回退 Normal/正文 样式定义
    if cjk_fonts:
        top_cjk = max(cjk_fonts, key=cjk_fonts.get)
        if _is_songti(top_cjk):
            items.append({"level": "ok", "msg": f"正文中文字体为「{top_cjk}」（宋体类，符合常见规范）"})
        else:
            items.append({"level": "warn",
                           "msg": f"正文中文字体为「{top_cjk}」，论文一般用宋体（以本校模板为准）"})
    else:
        style_ea = _style_east_asia_font(doc, set(BODY_STYLE_NAMES))
        if style_ea:
            if _is_songti(style_ea):
                items.append({"level": "ok",
                               "msg": f"正文中文字体由样式定义为「{style_ea}」（宋体类，符合常见规范）"})
            else:
                items.append({"level": "warn",
                               "msg": f"正文中文字体由样式定义为「{style_ea}」，论文一般用宋体（以本校模板为准）"})
        else:
            items.append({"level": "warn",
                           "msg": "正文中文字体未显式设置（继承自样式/主题，无法静态判定，建议人工确认）"})

    # 西文字体：Times New Roman 即符合常见规范
    if west_fonts:
        top_west = max(west_fonts, key=west_fonts.get)
        if "Times New Roman" in top_west:
            items.append({"level": "ok", "msg": f"正文西文字体为「{top_west}」（符合常见规范）"})
        else:
            items.append({"level": "warn",
                           "msg": f"正文西文字体为「{top_west}」，论文西文一般用 Times New Roman（以本校模板为准）"})

    if body_sizes:
        top_sz = max(body_sizes, key=body_sizes.get)
        if 10.5 <= top_sz <= 14:
            items.append({"level": "ok", "msg": f"正文主字号约 {top_sz}pt"})
        else:
            items.append({"level": "warn", "msg": f"正文主字号约 {top_sz}pt，常见为小四(12pt)/五号(10.5pt)"})
    elif not cjk_fonts and not west_fonts:
        items.append({"level": "warn", "msg": "正文未显式设置字号（继承自样式，无法静态判定，建议人工确认）"})

    # 摘要字数 / 关键词个数 / 中英文摘要
    paras = [p.text.strip() for p in doc.paragraphs]
    joined = "\n".join(paras)
    has_cn_abs = bool(re.search(r"^\s*摘\s*要\s*$", joined, re.M)) or "摘要" in joined[:4000]
    has_en_abs = bool(re.search(r"\bAbstract\b", joined, re.I))
    items.append({"level": "ok" if has_cn_abs else "warn",
                   "msg": "检出中文摘要" if has_cn_abs else "未检出中文摘要"})
    items.append({"level": "ok" if has_en_abs else "warn",
                   "msg": "检出英文摘要 Abstract" if has_en_abs else "未检出英文摘要 Abstract"})

    # 摘要正文字数（摘要标题到关键词之间）
    abs_start = next((i for i, t in enumerate(paras) if re.match(r"^\s*摘\s*要\s*$", t)), None)
    if abs_start is not None:
        chunk = []
        for t in paras[abs_start + 1: abs_start + 25]:
            if re.match(r"^\s*(关键词|关键字|key\s*words)", t, re.I):
                break
            chunk.append(t)
        n_char = sum(1 for ch in "".join(chunk) if not ch.isspace())
        if n_char == 0:
            items.append({"level": "warn", "msg": "中文摘要正文为空或无法定位"})
        elif n_char < 200:
            items.append({"level": "warn", "msg": f"中文摘要约 {n_char} 字，偏少（常见 300~600 字）"})
        elif n_char > 800:
            items.append({"level": "warn", "msg": f"中文摘要约 {n_char} 字，偏多（常见 300~600 字）"})
        else:
            items.append({"level": "ok", "msg": f"中文摘要约 {n_char} 字，长度适中"})

    kw_line = next((t for t in paras if re.match(r"^\s*(关键词|关键字)\s*[:：]", t)), None)
    if kw_line:
        kws = [k for k in re.split(r"[;；,，、]\s*", re.sub(r"^\s*(关键词|关键字)\s*[:：]", "", kw_line)) if k.strip()]
        if 3 <= len(kws) <= 8:
            items.append({"level": "ok", "msg": f"中文关键词 {len(kws)} 个（规范 3~8 个）"})
        else:
            items.append({"level": "warn", "msg": f"中文关键词 {len(kws)} 个，规范建议 3~8 个"})
    else:
        items.append({"level": "warn", "msg": "未检出“关键词：”行"})

    return items


# ============================================================
# 对外入口
# ============================================================

_TITLES = {
    "citation":         "引用对应性",
    "reference_format": "参考文献格式（GB/T 7714）",
    "structure":        "结构完整性",
    "compliance":       "格式合规",
}


def run_checks(thesis_path: str, review_path: str = None, categories=None) -> dict:
    """对论文执行所选检查。review_path 可为 docx/pdf/txt 文献综述文件。"""
    cats = [c for c in (categories or ALL_CATEGORIES) if c in ALL_CATEGORIES]
    if not cats:
        cats = ALL_CATEGORIES

    review_text = None
    if review_path:
        try:
            with open(review_path, "rb") as f:
                review_text, _ = extract_any(review_path, f.read())
        except Exception:
            review_text = None

    doc, paragraphs = _doc_paragraphs(thesis_path)

    result_cats = []
    for c in cats:
        if c == "citation":
            its = _check_citation(paragraphs, review_text)
        elif c == "reference_format":
            its = _check_reference_format(paragraphs)
        elif c == "structure":
            its = _check_structure(thesis_path)
        elif c == "compliance":
            its = _check_compliance(thesis_path)
        else:
            continue
        result_cats.append({"id": c, "title": _TITLES[c], "items": its})

    summary = {"ok": 0, "warn": 0, "error": 0}
    for rc in result_cats:
        for it in rc["items"]:
            summary[it["level"]] = summary.get(it["level"], 0) + 1

    return {"summary": summary, "categories": result_cats,
            "review_used": bool(review_text)}
